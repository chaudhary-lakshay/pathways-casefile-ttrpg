"""Convert the 23 published .xlsx character sheets into JSON fixtures.

These are FIXTURES, not content. Where a sheet disagrees with the
rulebook the sheet is wrong (spec section 6); we record what it printed
so the engine can show a computed-vs-published diff.

Reads the CharSheetP1 tab only. ThrustTable is content (Task 10), and
Summons is byte-identical boilerplate across all 23 files with no
character data at all.

The source directory is overridable (KNOWN_ISSUES: "Python tooling is not
reproducible off this machine") so this runs on any machine or CI without
the original spreadsheets, hardcoded only as the default for this
workstation:
  - `PATHWAYS_SHEETS_DIR` environment variable, or
  - `--src <dir>` on the command line (takes precedence over the env var).
"""
import argparse, glob, json, os, pathlib, re

ROOT = pathlib.Path(__file__).resolve().parent.parent
DEFAULT_SRC = pathlib.Path(os.path.expanduser("~/Downloads/LOTM TTRPG Sheets"))
OUT = ROOT / "fixtures"


def resolve_src_dir(cli_arg: str | None = None) -> pathlib.Path:
    if cli_arg:
        return pathlib.Path(cli_arg)
    env = os.environ.get("PATHWAYS_SHEETS_DIR")
    if env:
        return pathlib.Path(env)
    return DEFAULT_SRC

DIFFICULTIES = {"E", "A", "H", "VH"}


def slug(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_")


def num(v):
    """Cell -> int when integral, float when not, None when blank."""
    if v is None or v == "":
        return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return int(f) if f == int(f) else f


def find(grid, label):
    """Row/col of the first cell equal to `label`, or None."""
    for r, row in enumerate(grid):
        for c, cell in enumerate(row):
            if isinstance(cell, str) and cell.strip() == label:
                return r, c
    return None


def parse_skill(raw: str):
    """'History (Occult)/IQ' -> ('History', 'Occult', 'IQ')."""
    name_part, _, attr = raw.rpartition("/")
    m = re.match(r"^(.*?)\s*\((.*)\)\s*$", name_part.strip())
    if m:
        return m.group(1).strip(), m.group(2).strip(), attr.strip()
    return name_part.strip(), None, attr.strip()


def extract(path: pathlib.Path) -> dict:
    from openpyxl import load_workbook

    ws = load_workbook(str(path), data_only=True)["CharSheetP1"]
    grid = [list(r) for r in ws.iter_rows(values_only=True)]

    def value_right_of(label, offset=1):
        pos = find(grid, label)
        if pos is None:
            return None
        r, c = pos
        return num(grid[r][c + offset]) if c + offset < len(grid[r]) else None

    attributes = {a: value_right_of(a) for a in ("ST", "DX", "IQ", "HT")}

    def points_after(label):
        """First '<int> pts'/'<int> points'-shaped string cell after `label`, or None.

        Shared by "Budget (BP)"/"Total spent" (pre-existing) and the four
        POINT BREAKDOWN subtotal rows (Attributes/Advantages/Disadvantages/
        Skills) added for KNOWN_ISSUES: "fixtures carry no per-category
        point breakdown". Column position drifts by one for at least one
        subtotal row (Skills lands one column further right than
        Attributes/Advantages on every sheet observed), which is exactly
        why this scans forward for the "<n> pts" shape instead of reading
        a fixed offset. Monastery Scholar spells this row "30 points"
        rather than "30 pts" (every other sheet's Disadvantages row uses
        "pts"), so both spellings are accepted; a cell that is itself a
        broken formula (Church Healer's Disadvantages cell is the literal
        string "#ERROR!") correctly falls through to None rather than
        being misparsed.
        """
        pos = find(grid, label)
        if pos is None:
            return None
        r, c = pos
        for cell in grid[r][c + 1:]:
            if isinstance(cell, str):
                m = re.search(r"(\d+)\s*(?:pts|points)\b", cell)
                if m:
                    return int(m.group(1))
        return None

    published = {
        "hp":         value_right_of("HP"),
        "fp":         value_right_of("FP"),
        "will":       value_right_of("Will"),
        "per":        value_right_of("Per"),
        "spi":        value_right_of("SPI"),
        "cor":        value_right_of("CoR"),
        "basicSpeed": value_right_of("BASE SPD"),
        "basicMove":  value_right_of("BASE MOVE"),
        "dodge":      value_right_of("Dodge"),
        "bpTotal":    points_after("Budget (BP)"),
        "totalSpent": points_after("Total spent"),
        "unspent":    value_right_of("Unspent", 2),
        # POINT BREAKDOWN subtotals (Ch. 3/6). Beat Cop has no such block at
        # all, so every one of these is None for that sheet -- see
        # test_beat_cop_has_no_point_breakdown. Several other sheets HAVE a
        # block whose own printed subtotals contradict their own item lists
        # (documented for University Librarian, Street Acrobat, Monastery
        # Scholar in GoldenSheetTest) -- these are recorded as-published
        # here regardless; reconciling engine-vs-sheet is the golden
        # suite's job, not the extractor's.
        "attributesSubtotal":    points_after("Attributes"),
        "advantagesSubtotal":    points_after("Advantages"),
        "disadvantagesSubtotal": points_after("Disadvantages"),
        "skillsSubtotal":        points_after("Skills"),
    }

    def traits(header, subtotal, sheet_name):
        """Parse an ADVANTAGES/DISADVANTAGES block.

        Each block has its own two-row header ("ADVANTAGES AND PERKS" /
        "Name ... Roll Cost") immediately above the data rows. The cost
        column position is read off THAT header row rather than assumed,
        because some sheets also populate the "Roll" column for these
        entries (e.g. Fortune Teller's "Charisma 2", Street Acrobat's
        disadvantages) -- scanning for "the first number in the row"
        then grabs Roll instead of Cost.

        For most sheets the "Cost" column is simply correct. But at least
        two sheets (Fortune Teller advantages, Street Acrobat
        disadvantages) put the REAL cost in "Roll" instead, leaving "Cost"
        wrong for every row that has both populated -- reading "Cost"
        blindly silently ships bad data. The block's own printed subtotal
        (`subtotal`, from the "Advantages"/"Disadvantages" POINT BREAKDOWN
        row) is a checksum that disambiguates: when any row has both Roll
        and Cost populated, sum each column and use whichever sum equals
        `subtotal`. If neither reconciles (or there is no subtotal to
        check against), fall back to Cost -- the sheet is simply
        internally inconsistent (e.g. Monastery Scholar, University
        Librarian), and Cost remains the least-surprising default. This
        never fires for a block where Roll is blank throughout (e.g.
        Street Acrobat's single-row ADVANTAGES block), which always uses
        Cost exactly as before.
        """
        pos = find(grid, header)
        if pos is None:
            return []
        r0, c0 = pos
        header_row = grid[r0 + 1] if r0 + 1 < len(grid) else []
        cost_col = None
        roll_col = None
        for i, cell in enumerate(header_row[c0:]):
            if isinstance(cell, str):
                label = cell.strip()
                if label == "Cost" and cost_col is None:
                    cost_col = c0 + i
                elif label == "Roll" and roll_col is None:
                    roll_col = c0 + i
            if cost_col is not None and roll_col is not None:
                break
        rows = []
        for row in grid[r0 + 2:]:                    # +2 skips the sub-header row
            name = row[c0] if c0 < len(row) else None
            if not isinstance(name, str) or not name.strip():
                break
            if name.startswith("*GURPS"):
                break
            cost = num(row[cost_col]) if cost_col is not None and cost_col < len(row) else None
            roll = num(row[roll_col]) if roll_col is not None and roll_col < len(row) else None
            rows.append({"name": name.strip(), "cost": cost, "roll": roll})

        use_roll = False
        if subtotal is not None and any(r["cost"] is not None and r["roll"] is not None for r in rows):
            cost_sum = sum(r["cost"] or 0 for r in rows)
            roll_sum = sum(r["roll"] or 0 for r in rows)
            if roll_sum == subtotal and cost_sum != subtotal:
                use_roll = True
                print(
                    f"NOTICE: {sheet_name} {header}: Cost column sums to {cost_sum}, "
                    f"which does not match the printed subtotal ({subtotal}); using "
                    f"the Roll column instead, which sums to {roll_sum}."
                )

        out = []
        for r in rows:
            value = r["roll"] if use_roll else r["cost"]
            out.append({"name": r["name"], "cost": value or 0})
        return out

    skills = []
    pos = find(grid, "SKILLS & TECHNIQUE")
    if pos:
        r0, c0 = pos
        for row in grid[r0 + 2:]:
            cell = row[c0] if c0 < len(row) else None
            if not isinstance(cell, str) or not cell.strip():
                break
            text = cell.strip()
            if "/" not in text:                      # category header, e.g. "Social & Utility (56 pts):"
                continue
            rest = [x for x in row[c0 + 1:] if x is not None and x != ""]
            diff = next((str(x).strip() for x in rest if str(x).strip() in DIFFICULTIES), None)
            nums = [num(x) for x in rest if num(x) is not None]
            if diff is None or len(nums) < 2:
                continue
            name, spec, attr = parse_skill(text)
            skills.append({
                "name": name, "specialization": spec, "attr": attr,
                "difficulty": diff, "points": nums[1], "publishedLevel": nums[0],
            })

    return {
        "name": path.stem,
        "attributes": attributes,
        "published": published,
        "advantages": traits("ADVANTAGES AND PERKS", published["advantagesSubtotal"], path.stem),
        "disadvantages": traits("DISADVANTAGES AND QUIRKS", published["disadvantagesSubtotal"], path.stem),
        "skills": skills,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--src", default=None, help="directory of .xlsx sheets (overrides PATHWAYS_SHEETS_DIR)")
    args = parser.parse_args()
    src = resolve_src_dir(args.src)

    OUT.mkdir(parents=True, exist_ok=True)
    n = 0
    for f in sorted(glob.glob(str(src / "*.xlsx"))):
        p = pathlib.Path(f)
        if p.name.startswith("~$"):        # Excel lock file
            continue
        data = extract(p)
        (OUT / f"{slug(p.stem)}.json").write_text(
            json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
        n += 1
    print(f"{n} fixtures -> {OUT}")


if __name__ == "__main__":
    main()
